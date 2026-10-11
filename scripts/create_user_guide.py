from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".guide-deps"))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
import fitz


OUTPUT_DIR = ROOT / "delivery" / "인생흐름주기_의뢰인전달_20261002"
DOCX_PATH = OUTPUT_DIR / "인생흐름주기_사용안내서.docx"
PDF_PATH = OUTPUT_DIR / "인생흐름주기_사용안내서.pdf"
FOOTER = "의뢰인 전달용  |  v0.1.0  |  2026.10.02 수정"
INSTALLER_NAME = "인생흐름주기_설치파일.exe"
INPUT_IMAGE = ROOT / "artifacts" / "input-screen.png"
RESULT_IMAGE = ROOT / "artifacts" / "result-all.png"
PREVIEW_DIR = ROOT / "artifacts" / "user-guide-preview"

INK = RGBColor(44, 36, 26)
ACCENT = RGBColor(179, 51, 42)
MUTED = RGBColor(112, 96, 77)


def set_font(run, size=10.5, bold=False, color=INK):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_text(document, text, size=10.5, bold=False, color=INK, after=5, before=0):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    set_font(run, size, bold, color)
    return paragraph


def add_heading(document, text, level=2, before=8):
    size = 13 if level == 2 else 16
    color = ACCENT if level == 2 else INK
    paragraph = add_text(document, text, size=size, bold=True, color=color, before=before, after=5)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_bullet(document, parts):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.first_line_indent = Cm(-0.25)
    paragraph.paragraph_format.space_after = Pt(3)
    for text, bold in parts:
        set_font(paragraph.add_run(text), bold=bold)
    return paragraph


def add_screenshot(document, path, caption):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(path), width=Inches(6.25))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(5)
    set_font(caption_paragraph.add_run(caption), size=8.5, color=MUTED)


def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(4)
    set_font(header.add_run("인생 흐름 주기  |  사용 안내서"), size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run(FOOTER), size=8, color=MUTED)




def pdf_text(page, text, rect, size=10.5, bold=False, color=(0.17, 0.14, 0.10), align=0):
    return page.insert_textbox(
        fitz.Rect(*rect),
        text,
        fontname="MalgunBold" if bold else "Malgun",
        fontsize=size,
        lineheight=1.35,
        color=color,
        align=align,
    )


def prepare_pdf_page(pdf):
    page = pdf.new_page(width=595.28, height=841.89)
    page.insert_font(fontname="Malgun", fontfile=r"C:\Windows\Fonts\malgun.ttf")
    page.insert_font(fontname="MalgunBold", fontfile=r"C:\Windows\Fonts\malgunbd.ttf")
    pdf_text(page, "인생 흐름 주기  |  사용 안내서", (48, 30, 547, 48), size=8.5, color=(0.44, 0.38, 0.30))
    page.draw_line((48, 51), (547, 51), color=(0.81, 0.74, 0.62), width=0.7)
    pdf_text(page, FOOTER, (48, 805, 547, 822), size=8, color=(0.44, 0.38, 0.30), align=2)
    return page


B = True
N = False

# 각 항목은 (글자, 굵게) 조각의 목록이다. DOCX와 PDF가 같은 내용을 쓴다.
INSTALL = [
    [(INSTALLER_NAME, B), ("를 실행합니다.", N)],
    [("Windows 보안 경고가 나타나면 ", N), ("추가 정보", B), ("를 누른 뒤 ", N), ("실행", B), ("을 선택합니다.", N)],
    [("설치를 마친 뒤 바탕 화면의 ", N), ("인생 흐름 주기", B), ("를 실행합니다. 이전 버전이 설치되어 있어도 그대로 설치하면 새 버전으로 바뀝니다.", N)],
]

INPUT = [
    [("이름", B), (": 상담받는 고객의 이름을 입력합니다.", N)],
    [("생년", B), (": 고객이 1세가 되는 해를 입력합니다. 그래프의 1세는 항상 이 연도입니다.", N)],
    [("입춘나이", B), (": 숫자로 입력합니다. 0이나 음수(예: -16)도 입력할 수 있습니다.", N)],
    [("입력을 마치면 ", N), ("분석하기", B), ("를 누르거나 ", N), ("Enter", B), (" 키를 누릅니다.", N)],
]

RESULT = [
    [("그래프는 생년부터 약 100년의 흐름을 2.5년 간격으로 표시합니다.", N)],
    [("가로축 아래에는 ", N), ("5년마다 나이·연도·절기", B), ("가 표시되고, 그 사이 칸에는 절기만 표시됩니다.", N)],
    [("입춘나이가 그래프의 한 칸에 정확히 오도록 시작 나이(0세 또는 1세)가 자동으로 정해집니다. 자세한 내용은 5번을 참고하세요.", N)],
    [("심리, 수단, 재물, 명예 항목을 눌러 원하는 선을 표시하거나 숨길 수 있습니다.", N)],
    [("Ctrl + 마우스 휠", B), ("로 확대·축소하고, ", N), ("마우스 휠 또는 그래프 드래그", B), ("로 좌우 이동합니다.", N)],
    [("아래쪽 범위 막대의 양 끝은 확대·축소, 가운데는 구간 이동에 사용합니다.", N)],
    [("전체 보기", B), ("는 전체 구간을 표시하고, ", N), ("화면 맞춤", B), ("은 읽기 쉬운 구간으로 돌아갑니다.", N)],
]

PRINT = [
    [("오른쪽 위의 ", N), ("인쇄", B), ("를 누릅니다.", N)],
    [("종이 출력은 사용할 프린터와 ", N), ("A4 가로 방향", B), ("을 확인한 뒤 인쇄합니다.", N)],
    [("PDF 보관은 프린터에서 ", N), ("Microsoft Print to PDF", B), ("를 선택하고, 저장 위치와 파일명을 직접 지정합니다.", N)],
    [("결과지 위쪽에는 선행원 로고가, 아래쪽에는 상호와 연락처가 함께 인쇄됩니다.", N)],
]

NOTE = "정보 수정은 입력값을 유지한 채 이전 화면으로 돌아가고, 새 상담은 새 고객 입력 화면을 엽니다. 프로그램은 인터넷 연결 없이 사용할 수 있습니다."

CHANGES_INTRO = "2026년 10월 1일 밤부터 10월 2일 새벽까지 고친 내용입니다. 이번 설치 파일에 모두 들어 있습니다."

CHANGES = [
    ("입력 방식", [
        [("'1세 시작 절기' 선택 칸이 없어지고 ", N), ("'입춘나이' 숫자 입력 칸", B), ("으로 바뀌었습니다. 1세의 절기는 입춘나이로 자동 계산됩니다.", N)],
        [("생년은 언제나 1세가 되는 해입니다. 예를 들어 2003년생, 입춘나이 -16세를 넣으면 그래프의 1세는 2003년입니다.", N)],
        [("입춘나이 칸의 흐린 예시 글자는 1입니다.", N)],
    ]),
    ("그래프 시작 나이 자동 조정", [
        [("입춘나이 1, 3, 6, 8, 11, 13세 …와 0세 이하", B), (": 1세부터 시작합니다.", N)],
        [("입춘나이 2, 5, 7, 10, 12, 15세 …", B), (": 0세부터 시작하고, 0세 지점도 화면에 보입니다.", N)],
        [("입춘나이 4, 9, 14, 19세 …", B), (": 계산은 -1세부터 하고, 화면에는 1세부터 보입니다(1, 4, 6, 9세 …).", N)],
        [("이렇게 해서 입춘나이가 항상 그래프의 한 칸에 정확히 놓입니다.", N)],
    ]),
    ("그래프 표시", [
        [("가로축 아래 나이·연도 표시를 ", N), ("5년 단위", B), ("로 바꿨습니다. 그 사이 칸에는 절기만 표시됩니다.", N)],
        [("입춘 지점에 마우스를 올리면 입력한 입춘나이와 같은 나이가 나옵니다. 전에는 입춘나이 7세가 8세로 보였습니다.", N)],
        [("입춘나이 4, 9, 14세 …일 때 아래쪽 범위 막대가 오른쪽으로 치우쳐 보이던 문제를 고쳤습니다.", N)],
        [("결과 화면, 결과지, PNG 그림에서 ", N), ("'표시 시작 n세 · n년 절기'", B), (" 문구를 지웠습니다. 고객 정보에는 생년, 입춘나이, 상담일만 나옵니다.", N)],
    ]),
    ("선행원 상호 적용", [
        [("선행원 로고", B), ("를 상단 바, 바탕 화면·작업 표시줄 아이콘에 넣었습니다. 창 제목은 '선행원 · 인생 흐름 주기'입니다.", N)],
        [("입력 화면 맨 아래에 '선행원 행복을 찾아 주는 집 · 사주 · 타로 · 작명 · 궁합 · 심현 · 010-7579-2711' 줄을 넣었습니다.", N)],
        [("결과지 머리에 로고와 '행복을 찾아 주는 집'을, 꼬리말에 상호·상담 분야·연락처를 넣었습니다. 그래프 오른쪽 아래에는 작은 '선행원' 서명이 들어갑니다.", N)],
    ]),
    ("입력 화면 정리", [
        [("'상담 정보 입력' 제목 위의 '새 상담' 글자를 지웠습니다.", N)],
        [("입력 상자가 화면 정가운데에 오도록 위치를 맞췄습니다.", N)],
    ]),
]


def plain(parts):
    return "".join(text for text, _bold in parts)


class PdfWriter:
    LEFT, RIGHT, TOP, BOTTOM = 48, 547, 66, 790

    def __init__(self):
        self.pdf = fitz.open()
        self.new_page()

    def new_page(self):
        self.page = prepare_pdf_page(self.pdf)
        self.y = self.TOP

    def text(self, value, indent=0, size=10.5, bold=False, color=(0.17, 0.14, 0.10), gap=4, align=0):
        rect = (self.LEFT + indent, self.y, self.RIGHT - indent / 2, self.BOTTOM)
        left = pdf_text(self.page, value, rect, size=size, bold=bold, color=color, align=align)
        if left < 0:
            self.new_page()
            return self.text(value, indent, size, bold, color, gap, align)
        self.y = self.BOTTOM - left + gap

    def heading(self, value, size=13, color=(0.70, 0.20, 0.16)):
        if self.y > self.BOTTOM - 90:
            self.new_page()
        self.text(value, size=size, bold=True, color=color, gap=4)

    def bullets(self, items, indent=10):
        for parts in items:
            if self.y > self.BOTTOM - 30:
                self.new_page()
            pdf_text(self.page, "•", (self.LEFT + indent, self.y, self.LEFT + indent + 12, self.y + 20))
            self.text(plain(parts), indent=indent + 13, gap=2)
        self.y += 6

    def image(self, path, caption):
        if self.y + 310 > self.BOTTOM:
            self.new_page()
        rect = fitz.Rect(55, self.y + 2, 540, self.y + 287)
        self.page.insert_image(rect, filename=str(path))
        self.page.draw_rect(rect, color=(0.81, 0.74, 0.62), width=0.8)
        self.y = rect.y1 + 4
        self.text(caption, size=8.5, color=(0.44, 0.38, 0.30), align=1, gap=10)

    def note(self, title, body):
        if self.y + 80 > self.BOTTOM:
            self.new_page()
        top = self.y + 4
        rect = fitz.Rect(self.LEFT, top, self.RIGHT, top + 70)
        self.page.draw_rect(rect, color=(0.70, 0.20, 0.16), fill=(0.95, 0.93, 0.86), width=0.8)
        pdf_text(self.page, title, (60, top + 10, 535, top + 32), size=10.5, bold=True, color=(0.70, 0.20, 0.16))
        pdf_text(self.page, body, (60, top + 33, 535, top + 66), size=9.5)
        self.y = rect.y1 + 10


def add_pdf():
    w = PdfWriter()
    w.text("인생 흐름 주기", size=22, bold=True, gap=2)
    w.text("설치 및 사용 안내서", size=14, bold=True, color=(0.70, 0.20, 0.16), gap=10)
    w.heading("1. 설치하기")
    w.bullets(INSTALL)
    w.heading("2. 상담 정보 입력")
    w.bullets(INPUT)
    w.image(INPUT_IMAGE, "상담 정보 입력 화면")

    w.new_page()
    w.heading("3. 결과 확인", size=16, color=(0.17, 0.14, 0.10))
    w.bullets(RESULT)
    w.image(RESULT_IMAGE, "결과 그래프 화면 (1975년생, 입춘나이 -16세)")
    w.heading("4. 인쇄 또는 PDF로 보관")
    w.bullets(PRINT)
    w.note("알아두기", NOTE)

    w.new_page()
    w.heading("5. 수정된 부분", size=16, color=(0.17, 0.14, 0.10))
    w.text(CHANGES_INTRO, size=10, color=(0.44, 0.38, 0.30), gap=10)
    for title, items in CHANGES:
        w.heading(title, size=12)
        w.bullets(items)

    w.pdf.save(PDF_PATH, garbage=4, deflate=True)
    w.pdf.close()

    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    rendered = fitz.open(PDF_PATH)
    for index, page in enumerate(rendered):
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
        pixmap.save(PREVIEW_DIR / f"page-{index + 1}.png")
    rendered.close()


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_section(document.sections[0])

    normal = document.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK

    add_text(document, "인생 흐름 주기", size=22, bold=True, after=1)
    add_text(document, "설치 및 사용 안내서", size=14, bold=True, color=ACCENT, after=10)

    add_heading(document, "1. 설치하기", before=0)
    for parts in INSTALL:
        add_bullet(document, parts)

    add_heading(document, "2. 상담 정보 입력")
    for parts in INPUT:
        add_bullet(document, parts)
    add_screenshot(document, INPUT_IMAGE, "상담 정보 입력 화면")

    document.add_page_break()
    add_heading(document, "3. 결과 확인", level=1, before=0)
    for parts in RESULT:
        add_bullet(document, parts)
    add_screenshot(document, RESULT_IMAGE, "결과 그래프 화면 (1975년생, 입춘나이 -16세)")

    add_heading(document, "4. 인쇄 또는 PDF로 보관", before=6)
    for parts in PRINT:
        add_bullet(document, parts)

    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3ECDC")
    cell.margin_top = Cm(0.15)
    cell.margin_bottom = Cm(0.15)
    cell.margin_left = Cm(0.25)
    cell.margin_right = Cm(0.25)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    set_font(paragraph.add_run("알아두기\n"), bold=True, color=ACCENT)
    set_font(paragraph.add_run(NOTE), size=9.5)

    document.add_page_break()
    add_heading(document, "5. 수정된 부분", level=1, before=0)
    add_text(document, CHANGES_INTRO, size=10, color=MUTED, after=8)
    for title, items in CHANGES:
        add_text(document, title, size=12, bold=True, color=ACCENT, before=6, after=4).paragraph_format.keep_with_next = True
        for parts in items:
            add_bullet(document, parts)

    document.save(DOCX_PATH)
    add_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)


if __name__ == "__main__":
    main()
